"""Migration Check rule-pack support for the Migrate Studio CLI.

Migrate Studio

Copyright (c) 2026 Meta-Xceed, Inc.
SPDX-License-Identifier: Apache-2.0
"""

from __future__ import annotations

import json
import re
import shutil
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

# Locate the edition root and the default rule-pack spec.
ROOT = Path(__file__).resolve().parent.parent
DEFAULT_MC_SPEC = ROOT / "rules" / "migration_rules.json"
OUTPUT_DIRNAME = "Migrate_Studio"

# List the rule IDs that have an automatic update.
IMPLEMENTED_UPDATES = {"F1", "F2", "F3", "O1", "O2", "O3", "O5", "O6", "L3", "S1", "S5"}


# Describe one rule hit found while scanning a program line.
@dataclass
class Finding:
    category: str
    rule_id: str
    rule_name: str
    rule_desc: str
    criteria: str
    pgm_name: str
    pgm_path: str
    line_num: int
    line_text: str


# Describe one text change made by an automatic rule update.
@dataclass
class UpdateRecord:
    rule_id: str
    update_desc: str
    pgm_name: str
    pgm_path: str
    output_path: str
    line_num: int
    original_text: str
    updated_text: str


# Load the rule-pack JSON spec.
def load_spec(path: Path = DEFAULT_MC_SPEC) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


# Index rules by rule ID.
def rule_map(spec: dict) -> dict:
    return {r["rule_id"]: r for r in spec.get("rules", []) if r.get("rule_id")}


# Hide the internal S6 rule from reports.
def visible_rules(spec: dict) -> list[dict]:
    return [r for r in spec.get("rules", []) if r.get("rule_id") != "S6"]


# Strip a trailing parenthetical qualifier from a rule ID.
def clean_rule_id(rule_id: str) -> str:
    return re.sub(r"\(.*\)$", "", rule_id or "")


# Return implemented update rule IDs in the spec's application order.
def update_order(spec: dict) -> list[str]:
    order = spec.get("update_mode", {}).get("update_application_order", [])
    return [rid for rid in map(clean_rule_id, order) if rid in IMPLEMENTED_UPDATES]


# Return the line-level detection order from the spec.
def detection_order(spec: dict) -> list[str]:
    order = spec.get("rule_evaluation_order", {}).get("line_level_order", [])
    return [clean_rule_id(rid) for rid in order]


# Decode a SAS program as cp1252, preserving undecodable bytes.
def read_program(path: Path) -> str:
    return path.read_bytes().decode("cp1252", errors="surrogateescape")


# Encode text back to cp1252 without altering preserved bytes.
def write_program(path: Path, text: str) -> None:
    path.write_bytes(text.encode("cp1252", errors="surrogateescape"))


# List .sas files under a path, skipping output and archive folders.
def collect_programs(target: Path) -> list[Path]:
    if target.is_file():
        return [target] if target.suffix.lower() == ".sas" else []
    if not target.is_dir():
        raise SystemExit(f"No such program or folder: {target}")
    out = []
    for path in target.rglob("*.sas"):
        # Skip files inside generated output or archive directories.
        norm = "/" + path.as_posix().lower() + "/"
        if "/migrate_output/" in norm or "/migrate_studio/" in norm or "/archive/" in norm:
            continue
        out.append(path)
    return sorted(out)


# Detect a nearby WIN guard that exempts a line from Linux fixes.
def win_exempt(prev1: str, prev2: str, prev3: str) -> bool:
    return any('"WIN"' in p.upper() or "'WIN'" in p.upper() for p in (prev1, prev2, prev3))


# Classify a line as a Windows path that needs conversion.
def s5_classify(line: str) -> str:
    if "\\" not in line or re.match(r"^\s*/\*", line):
        return ""
    up = line.upper()
    if "PRXCHANGE" in up:
        return ""
    if re.search(r"[C-H]:\\", up) or "\\\\" in line or "..\\" in line or ".\\" in line:
        return "PATH"
    return ""


# Normalize line endings and split text into lines.
def split_lines(text: str) -> list[str]:
    return re.sub(r"\r\n?", "\n", text).split("\n")


# Build a Finding using rule metadata from the spec.
def mk_finding(spec: dict, rule_id: str, program: Path, root: Path, line_num: int,
               line_text: str, criteria: str) -> Finding:
    rule = rule_map(spec).get(rule_id, {})
    try:
        pgm_path = str(program.relative_to(root))
    except ValueError:
        pgm_path = str(program)
    return Finding(
        category=rule.get("category", ""),
        rule_id=rule_id,
        rule_name=rule.get("rule_name", rule_id),
        rule_desc=rule.get("description", ""),
        criteria=criteria,
        pgm_name=program.name,
        pgm_path=pgm_path,
        line_num=line_num,
        line_text=line_text,
    )


# Scan the whole program for chmod and %setup usage.
def whole_flags(lines: list[str], os_family: str) -> dict:
    flags = {"has_chmod": False, "has_setup_call": False}
    for line in lines:
        if re.search(r"\bchmod\b", line, re.I):
            flags["has_chmod"] = True
        if os_family == "LINUX":
            active = re.sub(r"/\*.*?\*/", "", line)
            if not re.match(r"^\s*\*", active) and re.search(r"%setup\b\s*(\(|;)", active, re.I):
                flags["has_setup_call"] = True
    return flags


# Read TOC file names from the spec, with defaults.
def configured_toc_filenames(spec: dict) -> list[str]:
    for parameter in spec.get("parameters", []):
        if parameter.get("name") == "TOC_FILENAMES":
            default = parameter.get("default", [])
            if isinstance(default, list):
                return [str(name) for name in default if str(name).strip()]
    return ["study_toc.xlsx", "toc.xlsx"]


# Map lowercase TOC file names to their on-disk casing.
def toc_case_map(root: Path, filenames: list[str] | None = None) -> dict:
    configured = {name.lower() for name in (filenames or ["study_toc.xlsx", "toc.xlsx"])}
    out = {}
    for path in sorted(root.rglob("*.xlsx")):
        low = path.name.lower()
        if low in configured and low not in out:
            out[low] = path.name
    return out


# Build a regex alternation that matches the TOC file names.
def toc_pattern(filenames: list[str]) -> str:
    return "(" + "|".join(re.escape(name) for name in filenames) + ")"


# Return the on-disk casing of a subfolder name, if it exists.
def find_case(parent: Path, ref: str) -> str | None:
    try:
        for item in parent.iterdir():
            if item.name.lower() == ref.lower() and item.is_dir():
                return item.name
    except OSError:
        return None
    return None


# Find path references whose folder casing differs from disk.
def f2_findings(line: str, program: Path, root: Path) -> list[tuple[str, str, str, int, int]]:
    if re.match(r"^\s*(%?\*|/\*)", line):
        return []
    out = []
    for match in re.finditer(r"[\\/ ]([^\\\/\"'\s;,()]+)", line):
        # Only path separators participate. The leading-space alternative is
        # not used; it keeps the regex scanner simple but is ignored here.
        if match.group(0)[0] not in "\\/":
            continue
        ref = match.group(1)
        if ref in ("", ".", ".."):
            continue
        # Extract the path text that precedes this reference.
        prefix = line[:match.start()]
        qpos = max(prefix.rfind('"'), prefix.rfind("'"))
        if qpos >= 0:
            token = prefix[qpos + 1:]
        else:
            token = prefix[max(prefix.rfind(ch) for ch in " =,();") + 1:]
        token = token.strip().replace("\\", "/").rstrip("/")
        if not token or "&" in token or "%" in token:
            continue
        # Resolve absolute paths directly and relative paths from the program folder.
        if re.match(r"^[A-Za-z]:", token) or token.startswith("//") or token.startswith("/"):
            parent = Path(token)
        else:
            parent = program.parent / token
        actual = find_case(parent, ref)
        if actual and actual != ref:
            out.append((ref, actual, str(parent), match.start() + 1, match.start() + 1 + len(ref)))
    return out


# Apply every detection rule to each line and return the findings.
def scan_program(spec: dict, text: str, program: Path, root: Path,
                 os_family: str = "LINUX") -> list[Finding]:
    lines = split_lines(text)
    flags = whole_flags(lines, os_family)
    findings: list[Finding] = []
    toc_names = configured_toc_filenames(spec)
    toc = toc_case_map(root, toc_names)
    toc_re = toc_pattern(toc_names)

    # Evaluate each line against the detection rules.
    for i, line in enumerate(lines):
        n = i + 1
        # Capture neighboring lines for context-sensitive rules.
        prev1 = lines[i - 1] if i > 0 else ""
        prev2 = lines[i - 2] if i > 1 else ""
        prev3 = lines[i - 3] if i > 2 else ""
        next1 = lines[i + 1] if i + 1 < len(lines) else ""
        next2 = lines[i + 2] if i + 2 < len(lines) else ""
        exempt = win_exempt(prev1, prev2, prev3)

        # Record a finding for the current line.
        def add(rid: str, criteria: str) -> None:
            findings.append(mk_finding(spec, rid, program, root, n, line, criteria))

        # Flag O1 (cd shell commands).
        if not exempt and (re.search(r"\bx\s*[\"']cd\s+", line, re.I) or
                           re.search(r"\bsystask\s+command\s*=\s*[\"']cd\s+", line, re.I)):
            add("O1", "Matched: cd shell command")
        # Flag O2 (%SYSGET(USERNAME)).
        if re.search(r"%sysget\s*\(\s*username\s*\)", line, re.I):
            add("O2", "Matched: %SYSGET(USERNAME)")
        # Flag O3 (filename _all_).
        if re.search(r"\bfilename\s+_all_\s*;", line, re.I):
            add("O3", "Matched: filename _all_;")
        # Flag S1 (_sasprogramfile).
        if re.search(r"_sasprogramfile", line, re.I):
            add("S1", "Matched: _sasprogramfile")
        # Flag S2 (proc append).
        if re.search(r"\bproc\s+append\b", line, re.I):
            add("S2", "Matched: proc append")
        # Flag S3 (Excel libname, export, or dbms=excel import).
        if re.search(r"\blibname\b", line, re.I) and re.search(r"\bexcel\b", line, re.I):
            add("S3", "Matched on same line: libname + excel")
        elif re.search(r"\bproc\s+export\b", line, re.I):
            add("S3", "Matched: proc export")
        elif re.search(r"\bproc\s+import\b", line, re.I) and re.search(r"\bdbms\s*=\s*excel\b", line, re.I):
            add("S3", "Matched on same line: proc import + dbms=excel")
        # Flag S4 (uppercase ERROR: and WARNING: tokens, including split-string variants).
        if any(token in line for token in ("ERROR:", "WARNING:")) or \
                ("ER" in line and "ROR:" in line) or ("WAR" in line and "NING:" in line):
            add("S4", "Matched uppercase ERROR:/WARNING: token")
        # Flag S5 (Windows path separators).
        if s5_classify(line) == "PATH":
            add("S5", "Matched S5 path indicator; Action=CONVERT")
        # Flag D1 (.bat references).
        if re.search(r"\.bat\b", line, re.I):
            add("D1", "Matched: .bat")
        # Flag D2 (.sh references with no chmod in the program).
        if re.search(r"\.sh(\s|[\"';]|$)", line, re.I) and not flags["has_chmod"]:
            add("D2", "Matched: .sh; Missing in program: chmod")
        # Flag D3 (smart quotes).
        if re.search("[\u2018\u2019]|\u00e2(?:\x80|\u20ac)(?:\x98|\x99|\u02dc|\u2122)", line):
            add("D3", "Matched: smart quote")
        # Flag C1 (source-platform folder names).
        if re.search(r"patient profiles[\\/]|[\\/]vendor|[\\/]profile|anadata[\\/]sdtm([\\/]|\\b)|reference[\\/]smq|validate[\\/]tlf[\\/]qcdata", line, re.I):
            add("C1", "Matched source-platform folder")
        # Flag C2 (sasuser).
        if re.search(r"\bsasuser\b", line, re.I):
            add("C2", "Matched: sasuser")
        # Flag L1 (standard Windows-side libnames).
        if not exempt and re.search(r"\blibname\s+(cdf|here|pfdata|vendordata|rawvend|raww|sdtma|smq|tlfout|tlfvdata|utility)\b", line, re.I):
            add("L1", "Matched standard Windows/source platform libname")
        # Flag L2 (filename pipe).
        if re.search(r"\bfilename\b", line, re.I) and re.search(r"\bpipe\b", line, re.I):
            add("L2", "Matched on same line: filename + pipe")
        # Flag L3 (del commands).
        if not re.match(r"^\s*(\*|/\*)", line) and re.search(r"\bx\s*[\"']del\s+|\bsystask\s+command\s*=\s*\"del\s+", line, re.I):
            add("L3", "Matched: x/systask del command")
        # Flag L4 (non-ASCII characters that risk encoding problems).
        if any(ord(ch) > 127 for ch in line):
            risky = [f"pos{i + 1}=[HEX:{ord(ch):02X}]" for i, ch in enumerate(line) if ord(ch) > 127]
            add("L4", f"Encoding risk count={len(risky)}; " + ", ".join(risky))
        # Flag O5 (xlsx proc import with no validvarname=v7 before it).
        o5_context = " ".join(part for part in (line, next1, next2) if part.strip())
        if not re.match(r"^\s*(\*|/\*)", line) and re.search(r"\bproc\s+import\b", line, re.I) and \
                (re.search(r"\.xlsx\b", o5_context, re.I) or re.search(r"\bdbms\s*=\s*xlsx\b", o5_context, re.I)) and \
                not re.match(r"^\s*options\s+validvarname\s*=\s*v7\s*;", prev1, re.I):
            add("O5", "Matched: proc import plus .xlsx or dbms=xlsx")
        # Flag O6 (copy commands).
        if re.search(r"\bx\s*\"copy\s+", line, re.I) or re.search(r"\bsystask\s+command\s*=\s*\"cmd\s+/c\s+copy\s+", line, re.I):
            add("O6", "Matched: x \"copy or systask command=\"cmd /c copy")
        # Flag F3 (xsync, xwait, and xcmd options).
        if not exempt and re.search(r"\boptions\b", line, re.I):
            options = re.findall(r"\b(noxsync|xsync|noxwait|xwait|noxcmd|xcmd)\b", line, re.I)
            if options:
                add("F3", " | ".join(f"Option=[{opt}]" for opt in options))
        # Flag F1 (TOC file names whose casing differs from disk).
        for match in re.finditer(toc_re, line, re.I):
            expected = toc.get(match.group(1).lower(), "")
            if expected and match.group(1) != expected:
                add("F1", f"Referenced=[{match.group(1)}] Physical file=[{expected}]")
        # Flag F2 (folder names whose casing differs from disk).
        for ref, actual, parent, *_ in f2_findings(line, program, root):
            add("F2", f"Referenced=[{ref}] Physical=[{actual}] Parent=[{parent}]")

    # Flag O4 once per Linux program that never calls %setup.
    if os_family == "LINUX" and not flags["has_setup_call"]:
        findings.append(mk_finding(spec, "O4", program, root, 0, "Program-level check: missing standard %setup macro invocation.",
                                   "Linux only: missing standard %setup macro invocation"))
    return findings


# Look up the update description for a rule.
def update_desc(spec: dict, rule_id: str) -> str:
    return rule_map(spec).get(rule_id, {}).get("update", {}).get("update_desc", "")


# Apply a per-line function and collect its output lines and records.
def map_lines(text: str, fn):
    out, records = [], []
    lines = split_lines(text)
    for i, line in enumerate(lines):
        new_lines, rec = fn(line, i, lines)
        out.extend(new_lines)
        records.extend(rec)
    return "\n".join(out), records


# Return a program's name and root-relative path.
def scan_ctx(program: Path, root: Path) -> tuple[str, str]:
    try:
        rel = str(program.relative_to(root))
    except ValueError:
        rel = str(program)
    return program.name, rel


# Build an update record for one rule change.
def rec(spec: dict, rule_id: str, program: Path, root: Path, line_num: int,
        before: str, after: str, output_path: str = "") -> UpdateRecord:
    pgm_name, pgm_path = scan_ctx(program, root)
    return UpdateRecord(rule_id, update_desc(spec, rule_id), pgm_name, pgm_path,
                        output_path, line_num, before, after)


# Apply the selected rule updates in spec order and return text and records.
def update_program(spec: dict, text: str, program: Path, root: Path,
                   selected: set[str] | None = None) -> tuple[str, list[UpdateRecord]]:
    records: list[UpdateRecord] = []
    toc_names = configured_toc_filenames(spec)
    toc = toc_case_map(root, toc_names)
    toc_re = toc_pattern(toc_names)
    selected = selected or IMPLEMENTED_UPDATES

    for rule_id in update_order(spec):
        if rule_id not in selected:
            continue
        if rule_id == "F1":
            # Rewrite TOC file names to match their on-disk casing (F1).
            def repl(match):
                return toc.get(match.group(1).lower()) or match.group(1)
            before = text
            text = re.sub(toc_re, repl, text, flags=re.I)
            if text != before:
                records.append(rec(spec, rule_id, program, root, 0, before, text))
        elif rule_id == "F2":
            # Fix folder-name casing in path references, up to 50 passes per line (F2).
            def fix_line(line, i, lines):
                next_line = line
                changed = False
                for _ in range(50):
                    found = f2_findings(next_line, program, root)
                    if not found:
                        break
                    ref, actual, parent, start, end = found[0]
                    next_line = next_line[:start] + actual + next_line[end:]
                    changed = True
                return [next_line], [rec(spec, rule_id, program, root, i + 1, line, next_line)] if changed else []
            text, made = map_lines(text, fix_line)
            records.extend(made)
        elif rule_id == "F3":
            # Prefix Windows-only options statements with a remove-for-linux marker (F3).
            def fix_line(line, i, lines):
                if re.match(r"^\s*\*{3}\s*remove for linux\s*\*{3}", line, re.I):
                    return [line], []
                prev = [lines[i - 1] if i > 0 else "", lines[i - 2] if i > 1 else "", lines[i - 3] if i > 2 else ""]
                if not win_exempt(*prev) and re.search(r"\boptions\b", line, re.I) and \
                        re.search(r"\b(noxsync|xsync|noxwait|xwait|noxcmd|xcmd)\b", line, re.I):
                    new = "*** remove for linux *** " + line
                    return [new], [rec(spec, rule_id, program, root, i + 1, line, new)]
                return [line], []
            text, made = map_lines(text, fix_line)
            records.extend(made)
        elif rule_id == "O1":
            # Replace cd commands with dlgcdir calls (O1).
            text, made = _update_o1(spec, text, program, root)
            records.extend(made)
        elif rule_id == "O2":
            # Replace %SYSGET(USERNAME) with %SYSGET(USER) (O2).
            before = text
            text = re.sub(r"%sysget\s*\(\s*username\s*\)", "%sysget(USER)", text, flags=re.I)
            if text != before:
                records.append(rec(spec, rule_id, program, root, 0, before, text))
        elif rule_id == "O3":
            # Preserve the _GSFNAME fileref around filename _all_ (O3).
            text, made = _update_o3(spec, text, program, root)
            records.extend(made)
        elif rule_id == "O5":
            # Add options validvarname=v7 before xlsx imports (O5).
            text, made = _update_o5(spec, text, program, root)
            records.extend(made)
        elif rule_id == "O6":
            # Replace copy commands with fcopy steps (O6).
            text, made = _update_o6(spec, text, program, root)
            records.extend(made)
        elif rule_id == "L3":
            # Replace del commands with rm -f (L3).
            text, made = _update_l3(spec, text, program, root)
            records.extend(made)
        elif rule_id == "S1":
            # Replace &_SASPROGRAMFILE with %SYSGET(SAS_EXECFILEPATH) (S1).
            before = text
            text = re.sub(r"&_SASPROGRAMFILE", "%sysget(SAS_EXECFILEPATH)", text, flags=re.I)
            if text != before:
                records.append(rec(spec, rule_id, program, root, 0, before, text))
        elif rule_id == "S5":
            # Convert backslashes to forward slashes on Windows path lines (S5).
            def fix_line(line, i, lines):
                if s5_classify(line) != "PATH":
                    return [line], []
                new = line.replace("\\", "/")
                return [new], [rec(spec, rule_id, program, root, i + 1, line, new)]
            text, made = map_lines(text, fix_line)
            records.extend(made)
    return text, records


# Replace cd commands with dlgcdir calls and keep the original as a note.
def _update_o1(spec, text, program, root):
    def fix_line(line, i, lines):
        # Skip lines that were already updated.
        if re.match(r"^\s*\*{3}\s*Updated to use dlgcdir\(\)\s*\*{3}", line, re.I):
            return [line], []
        prev = [lines[i - 1] if i > 0 else "", lines[i - 2] if i > 1 else "", lines[i - 3] if i > 2 else ""]
        # Skip lines guarded by a WIN check.
        if win_exempt(*prev):
            return [line], []
        # Extract the directory from a systask or x cd command.
        match = re.search(r"\bsystask\s+command\s*=\s*[\"']cd\s+([^\"']+)[\"']", line, re.I)
        if not match:
            match = re.search(r"\bx\s*[\"']cd\s+([^\"']+)[\"']", line, re.I)
        if not match:
            return [line], []
        # Emit a dlgcdir step that reports its return code.
        path = match.group(1).replace("\\", "/")
        new = ["data _null_;", f'   rc = dlgcdir("{path}");',
               '   put "NOTE: Change directory return code (0 means success): " rc;',
               "run;", "*** Updated to use dlgcdir() *** " + line]
        return new, [rec(spec, "O1", program, root, i + 1, line, "\n".join(new))]
    return map_lines(text, fix_line)


# Wrap filename _all_ so the _GSFNAME fileref is saved and restored.
def _update_o3(spec, text, program, root):
    # Detect the marker left by a prior update.
    marker = re.compile(r"^\s*\*{3}\s*Updated to preserve _GSFNAME\s*\*{3}\s*;", re.I)
    def fix_line(line, i, lines):
        if not re.search(r"\bfilename\s+_all_\s*;", line, re.I):
            return [line], []
        # Skip statements whose preceding lines already carry the marker.
        if any(marker.search(prev or "") for prev in lines[max(0, i - 8):i]):
            return [line], []
        # Emit the save step, the original statement, and the restore step.
        new = ["*** Updated to preserve _GSFNAME ***;", "%let old_path=;", "",
               "data _null_;", "    if fileref('_GSFNAME') = 0 then",
               "    call symputx('old_path', pathname('_GSFNAME'), 'g');", "run;",
               line, "data _null_;",
               "    if symget('old_path') ne '' then call execute(cats('filename _GSFNAME \"', symget('old_path'), '\";'));",
               "    else call execute('filename _GSFNAME temp;');", "run;"]
        return new, [rec(spec, "O3", program, root, i + 1, line, "\n".join(new))]
    return map_lines(text, fix_line)


# Insert options validvarname=v7 before xlsx proc import statements.
def _update_o5(spec, text, program, root):
    def fix_line(line, i, lines):
        prev = lines[i - 1] if i > 0 else ""
        next1 = lines[i + 1] if i + 1 < len(lines) else ""
        next2 = lines[i + 2] if i + 2 < len(lines) else ""
        ctx = " ".join(part for part in (line, next1, next2) if part.strip())
        # Skip comments and lines that are not proc import.
        if re.match(r"^\s*(\*|/\*)", line) or not re.search(r"\bproc\s+import\b", line, re.I):
            return [line], []
        # Require an xlsx file or dbms=xlsx in the statement.
        if not (re.search(r"\.xlsx\b", ctx, re.I) or re.search(r"\bdbms\s*=\s*xlsx\b", ctx, re.I)):
            return [line], []
        # Skip imports already preceded by the option.
        if re.match(r"^\s*options\s+validvarname\s*=\s*v7\s*;", prev, re.I):
            return [line], []
        new = ["options validvarname=v7;", line]
        return new, [rec(spec, "O5", program, root, i + 1, line, "\n".join(new))]
    return map_lines(text, fix_line)


# Parse the source and destination from an x copy command.
def _copy_payload(line: str) -> tuple[str, str] | None:
    match = re.match(r"^\s*x\s*[\"']copy\s+(.*)[\"']\s*;?\s*$", line, re.I)
    if not match:
        return None
    payload = match.group(1).strip()
    quoted = re.match(r'^""([^"]+)""\s+""([^"]+)""', payload)
    if quoted:
        return quoted.group(1), quoted.group(2)
    parts = payload.replace('""', '"').split()
    if len(parts) < 2:
        return None
    return parts[0].strip('"'), parts[1].strip('"')


# Replace x copy commands with an fcopy data step.
def _update_o6(spec, text, program, root):
    def fix_line(line, i, lines):
        if re.match(r"^\s*\*{3}\s*Update to use fcopy \( \)\s*\*{3}", line, re.I):
            return [line], []
        pair = _copy_payload(line)
        if not pair:
            return [line], []
        src, dest = pair
        # Emit filenames, an fcopy call, and a status note.
        new = ["*** Update to use fcopy ( ) *** " + line, f'filename src  "{src}";',
               f'filename dest "{dest}";', "", "data _null_;", "    rc = fcopy('src', 'dest');",
               f'    if rc = 0 then put "NOTE: Copied {src} to {dest}";',
               '    else put "ERROR: Copy failed. " rc= sysmsg();', "run;"]
        return new, [rec(spec, "O6", program, root, i + 1, line, "\n".join(new))]
    return map_lines(text, fix_line)


# Replace Windows del commands with rm -f.
def _update_l3(spec, text, program, root):
    def fix_line(line, i, lines):
        if re.match(r"^\s*\*{3}\s*Update del command with rm\s*\*{3}\s*;", line, re.I):
            return [line], []
        if re.match(r"^\s*(\*|/\*)", line):
            return [line], []
        # Rewrite x del commands first, then systask del commands.
        new = re.sub(r"(\bx\s*[\"'])del\s+", r"\1rm -f ", line, count=1, flags=re.I)
        if new == line:
            new = re.sub(r"(\bsystask\s+command\s*=\s*\")del\s+", r"\1rm -f ", line, count=1, flags=re.I)
        if new == line:
            return [line], []
        out = "*** Update del command with rm ***; " + new
        return [out], [rec(spec, "L3", program, root, i + 1, line, out)]
    return map_lines(text, fix_line)


# Copy a program into a dated archive folder and validate the copy.
def archive_backup(path: Path) -> Path:
    archive = path.parent / "archive"
    archive.mkdir(exist_ok=True)
    # Pick a unique archive name from the date and a counter.
    stamp = datetime.now().strftime("%Y%m%d")
    base = archive / f"{path.stem}_{stamp}{path.suffix}"
    backup = base
    i = 2
    while backup.exists():
        backup = archive / f"{path.stem}_{stamp}_{i}{path.suffix}"
        i += 1
    # Copy with metadata and verify size and timestamp.
    shutil.copy2(path, backup)
    src, dst = path.stat(), backup.stat()
    if dst.st_size == 0 or src.st_size != dst.st_size or src.st_mtime_ns != dst.st_mtime_ns:
        raise RuntimeError(f"Backup validation failed for {path}")
    return backup


# Write findings, updates, and rule metadata to an Excel report.
def write_workbook(path: Path, findings: list[Finding], updates: list[UpdateRecord],
                   spec: dict, os_family: str, root: Path) -> None:
    import openpyxl

    # Create the details sheet first.
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Migration Check Details"
    ws.append(["Category", "Rule ID", "Rule Name", "Rule Description", "Criteria",
               "Program Name", "Program Path", "Line Number", "Line Text"])
    for f in findings:
        if f.rule_id != "S7":
            ws.append([f.category, f.rule_id, f.rule_name, f.rule_desc, f.criteria,
                       f.pgm_name, f.pgm_path, f.line_num, f.line_text])

    # Add a placeholder sheet for global macro details.
    ws = wb.create_sheet("Global Macro Details")
    ws.append(["Rule ID", "Details"])
    ws.append(["S7", "Not evaluated - SASAUTOS support is not available in this CLI path yet."])

    # Add the rules list sheet.
    ws = wb.create_sheet("Rules List")
    ws.append(["Category", "Rule ID", "Rule Name", "Rule Description", "MC Special Rule", "Criteria"])
    for r in visible_rules(spec):
        ws.append([r.get("category", ""), r.get("rule_id", ""), r.get("rule_name", ""),
                   r.get("description", ""), r.get("special_rule", ""), r.get("criteria_template", "")])

    # Add a per-program, per-rule finding count sheet.
    ws = wb.create_sheet("Migration Check Summary")
    ws.append(["Program Path", "Rule ID", "Rule Name", "Findings"])
    counts = {}
    for f in findings:
        key = (f.pgm_path, f.rule_id, f.rule_name)
        counts[key] = counts.get(key, 0) + 1
    for (pgm_path, rule_id, rule_name), count in sorted(counts.items()):
        ws.append([pgm_path, rule_id, rule_name, count])

    # Add the update details sheet.
    ws = wb.create_sheet("Update Details")
    ws.append(["Rule ID", "Update Description", "Program Name", "Program Path",
               "Output Path", "Line Number", "Original Text", "Updated Text"])
    for u in updates:
        ws.append([u.rule_id, u.update_desc, u.pgm_name, u.pgm_path, u.output_path,
                   u.line_num, u.original_text, u.updated_text])

    # Add report metadata.
    ws = wb.create_sheet("Report Information")
    ws.append(["Field", "Value"])
    ws.append(["MC spec", str(DEFAULT_MC_SPEC)])
    ws.append(["Macro version", spec.get("source", {}).get("macro_version", "")])
    ws.append(["OS family", os_family])
    ws.append(["Root", str(root)])
    ws.append(["Generated", datetime.now().isoformat(timespec="seconds")])

    # Create the output folder and save the workbook.
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
